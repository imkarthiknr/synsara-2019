import csv
import io

from openpyxl import load_workbook

from app.content import get_content
from app.schemas import EventRegistration
from app.services import registrations
from tests.conftest import PDF, csrf_of, login, person, team


def seed(db, **overrides):
    return registrations.register(
        db, get_content(), EventRegistration.model_validate(person(**overrides))
    )


def test_admin_requires_login(client):
    for path in ["/admin", "/admin/participants", "/admin/teams", "/admin/export.xlsx"]:
        r = client.get(path, follow_redirects=False)
        assert r.status_code == 303, path
        assert r.headers["location"].startswith("/admin/login?next=/admin")


def test_wrong_password(client):
    r = login(client, password="nope")
    assert r.status_code == 401
    assert client.get("/admin", follow_redirects=False).status_code == 303


def test_login_ignores_external_next(client):
    token = csrf_of(client, "/admin/login")
    r = client.post(
        "/admin/login",
        data={
            "csrf": token,
            "username": "admin",
            "password": "test-admin-password",
            "next": "https://evil.example/",
        },
        follow_redirects=False,
    )
    assert r.headers["location"] == "/admin"


def test_dashboard_and_lists(client, db):
    seed(db)
    assert login(client).status_code == 303
    r = client.get("/admin")
    assert r.status_code == 200 and "Code Fest" in r.text
    assert "Asha Raman" in client.get("/admin/participants?event=code-fest").text
    assert "Asha Raman" not in client.get("/admin/participants?event=code-relay").text


def test_logout(client):
    login(client)
    token = csrf_of(client, "/admin")
    client.post("/admin/logout", data={"csrf": token})
    assert client.get("/admin", follow_redirects=False).status_code == 303


def test_delete_participant_requires_csrf(client, db):
    p = seed(db)
    login(client)
    assert client.post(f"/admin/participants/{p.id}/delete", data={"csrf": "x"}).status_code == 400
    token = csrf_of(client, "/admin")
    client.post(f"/admin/participants/{p.id}/delete", data={"csrf": token})
    assert registrations.list_participants(db) == []


def test_excel_export_and_formula_guard(client, db):
    seed(db, name="=HYPERLINK(1)", college="+cmd|calc")
    login(client)
    r = client.get("/admin/export.xlsx")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "All participants"
    assert "Code Fest" in wb.sheetnames
    assert "HACK-una-Matata" in " ".join(wb.sheetnames)
    row = [c.value for c in wb["All participants"][2]]
    assert "'=HYPERLINK(1)" in row and "'+cmd|calc" in row
    assert wb["Code Fest"].max_row == 2
    assert client.get("/admin/export/code-fest.xlsx").status_code == 200
    assert client.get("/admin/export/nope.xlsx").status_code == 404


def test_csv_export(client, db):
    seed(db)
    login(client)
    r = client.get("/admin/participants.csv")
    rows = list(csv.reader(io.StringIO(r.text.lstrip("﻿"))))
    assert rows[1][0] == "SYN19-0001"


def test_abstract_download(client):
    token = csrf_of(client, "/hackathon/register")
    client.post(
        "/hackathon/register",
        data={"csrf": token, **team()},
        files={"abstract": ("idea.pdf", PDF, "application/pdf")},
    )
    login(client)
    assert "Byte Busters" in client.get("/admin/teams").text
    r = client.get("/admin/teams/1/abstract")
    assert r.status_code == 200
    assert r.content == PDF
    assert "HUM19-001" in r.headers["content-disposition"]
    assert client.get("/admin/teams/99").status_code == 404


def test_demo_seed_is_idempotent(db):
    from app.seed import seed

    people, teams = seed()
    assert people == 12 and teams == 2
    assert seed() == (0, 0)
    assert registrations.counts_by_event(db)["code-fest"] == 5
