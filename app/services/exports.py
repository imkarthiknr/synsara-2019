"""Excel and CSV exports for organisers (the 2019 site wrote one .xlsx per event)."""

import csv
import io
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.content import Content
from app.models import HackathonTeam, Participant

PARTICIPANT_HEADERS = [
    "Reg. ID",
    "Name",
    "College",
    "Department",
    "Year",
    "Email",
    "Mobile",
    "Events",
    "Registered at",
]
TEAM_HEADERS = [
    "Team ID",
    "Team",
    "Leader",
    "Members",
    "College",
    "Department",
    "Year",
    "Email",
    "Mobile",
    "Project",
    "Domain",
    "Tech help",
    "Accommodation",
    "Heard from",
    "Abstract",
    "Registered at",
]
HEADER_FILL = PatternFill("solid", fgColor="0B0F0B")
HEADER_FONT = Font(bold=True, color="39FF14")


def _sheet_title(name: str) -> str:
    # Excel: max 31 chars, no []:*?/\
    return "".join(c for c in name if c not in "[]:*?/\\")[:31]


def _safe(value: object) -> object:
    """Neutralise spreadsheet formula injection in user-supplied text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def participant_row(content: Content, p: Participant) -> list[object]:
    names = [
        content.event(e.event_slug).name if content.event(e.event_slug) else e.event_slug
        for e in p.entries
    ]
    return [
        p.code,
        p.name,
        p.college,
        p.department,
        p.year,
        p.email,
        p.mobile,
        ", ".join(names),
        p.created_at.strftime("%Y-%m-%d %H:%M") if p.created_at else "",
    ]


def team_row(t: HackathonTeam) -> list[object]:
    return [
        t.code,
        t.team_name,
        t.leader_name,
        ", ".join(m.name for m in t.members),
        t.college,
        t.department,
        t.year,
        t.email,
        t.mobile,
        t.project_title,
        t.domain,
        "Yes" if t.needs_technical_help else "No",
        "Yes" if t.needs_accommodation else "No",
        t.heard_from,
        t.abstract_filename,
        t.created_at.strftime("%Y-%m-%d %H:%M") if t.created_at else "",
    ]


def _write_sheet(ws, headers: list[str], rows: Iterable[list[object]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    widths = [len(h) for h in headers]
    for row in rows:
        ws.append([_safe(v) for v in row])
        widths = [max(w, len(str(v))) for w, v in zip(widths, row, strict=False)]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 48)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def workbook(
    content: Content, participants: list[Participant], teams: list[HackathonTeam]
) -> bytes:
    """One workbook: 'All participants', one sheet per event, and the hackathon teams."""
    wb = Workbook()
    ws = wb.active
    ws.title = "All participants"
    _write_sheet(ws, PARTICIPANT_HEADERS, (participant_row(content, p) for p in participants))
    for event in content.events:
        rows = [
            participant_row(content, p)
            for p in participants
            if any(e.event_slug == event.slug for e in p.entries)
        ]
        _write_sheet(wb.create_sheet(_sheet_title(event.name)), PARTICIPANT_HEADERS, rows)
    _write_sheet(
        wb.create_sheet(_sheet_title(content.hackathon.name)),
        TEAM_HEADERS,
        (team_row(t) for t in teams),
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def event_workbook(content: Content, event_name: str, participants: list[Participant]) -> bytes:
    wb = Workbook()
    _write_sheet(
        wb.active, PARTICIPANT_HEADERS, (participant_row(content, p) for p in participants)
    )
    wb.active.title = _sheet_title(event_name)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def participants_csv(content: Content, participants: list[Participant]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(PARTICIPANT_HEADERS)
    for p in participants:
        w.writerow([_safe(v) for v in participant_row(content, p)])
    return buf.getvalue()
